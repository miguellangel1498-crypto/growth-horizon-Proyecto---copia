from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ENCABEZADO_FILL = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
ENCABEZADO_FONT = Font(color="FFFFFF", bold=True)
TITULO_FONT = Font(size=14, bold=True, color="0F766E")


def _ajustar_columnas(ws):
    for col in ws.columns:
        letra = get_column_letter(col[0].column)
        ancho = max((len(str(c.value or "")) for c in col), default=8) + 2
        ws.column_dimensions[letra].width = min(ancho, 60)


def _crear_libro(titulo, hoja, encabezados, filas):
    wb = Workbook()
    ws = wb.active
    ws.title = hoja[:31]

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(encabezados))
    celda_titulo = ws.cell(row=1, column=1, value=titulo)
    celda_titulo.font = TITULO_FONT

    fila_encabezados = 3
    for idx, encabezado in enumerate(encabezados, start=1):
        c = ws.cell(row=fila_encabezados, column=idx, value=encabezado)
        c.fill = ENCABEZADO_FILL
        c.font = ENCABEZADO_FONT
        c.alignment = Alignment(horizontal="center")

    for fila in filas:
        ws.append(fila)

    _ajustar_columnas(ws)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def exportar_diagnostico(filas):
    encabezados = ["Empresa", "Sector", "Tamano", "Indice General", "Segmento", "Fecha"]
    return _crear_libro(
        "Reporte de Diagnostico - Growth Horizon",
        "Diagnostico",
        encabezados,
        [list(fila) for fila in filas],
    )
